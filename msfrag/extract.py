"""Text extraction: turns a file into pages of plain text.

Every extractor returns an `Extracted` with a list of (page_number, text). "Page" means a PDF page, a slide,
a spreadsheet sheet, or 1 for formats without pages. Headings are emitted as Markdown "# " lines so the
chunker can track the section each chunk belongs to.
"""
import csv
import email
import email.policy
import io
import json
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".log", ".ini", ".cfg", ".conf", ".toml", ".yaml", ".yml",
            ".xml", ".sql", ".py", ".js", ".ts", ".java", ".c", ".h", ".cpp", ".cs", ".go", ".rs", ".sh",
            ".ps1", ".bat", ".php", ".rb", ".css", ".tex"}


@dataclass
class Extracted:
    kind: str                                   # pdf, docx, text, ...
    pages: list = field(default_factory=list)   # [(page_no, text)]
    title: str | None = None
    meta: dict = field(default_factory=dict)
    warning: str | None = None

    @property
    def chars(self):
        return sum(len(t) for _, t in self.pages)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    try:
        return data.decode("cp1252")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ---- per format --------------------------------------------------------------------------
def _text(path: Path):
    text = _clean(_decode(path.read_bytes()))
    kind = "markdown" if path.suffix.lower() in (".md", ".markdown") else "text"
    title = next((m.group(1).strip() for m in [re.search(r"^#\s+(.+)$", text, re.M)] if m), None)
    return Extracted(kind, [(1, text)], title)


def _pdf(path: Path):
    from pypdf import PdfReader
    r = PdfReader(str(path))
    pages = []
    for i, p in enumerate(r.pages, 1):
        try:
            pages.append((i, _clean(p.extract_text() or "")))
        except Exception:   # one broken page should not lose the document
            pages.append((i, ""))
    info = r.metadata or {}
    title = (info.get("/Title") or "").strip() or None
    meta = {k: str(info.get(f"/{k}")) for k in ("Author", "Subject", "Creator", "CreationDate") if info.get(f"/{k}")}
    meta["page_count"] = len(r.pages)
    ex = Extracted("pdf", [p for p in pages if p[1]], title, meta)
    if len(r.pages) and ex.chars < 40 * len(r.pages):
        ex.warning = "little or no text layer (scanned PDF?) - OCR is not supported yet"
    return ex


def _docx(path: Path):
    import docx
    d = docx.Document(str(path))
    lines = []
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = docx.text.paragraph.Paragraph(block, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"Heading (\d)", style)
            if m or style == "Title":
                lines.append("#" * (int(m.group(1)) if m else 1) + " " + text)
            elif style.startswith("List"):
                lines.append("- " + text)
            else:
                lines.append(text)
        elif tag == "tbl":
            t = docx.table.Table(block, d)
            for row in t.rows:
                cells = []
                for c in row.cells:
                    v = c.text.strip().replace("\n", " ")
                    if not cells or cells[-1] != v:   # merged cells repeat their text
                        cells.append(v)
                lines.append("| " + " | ".join(cells) + " |")
        lines.append("")
    cp = d.core_properties
    meta = {k: str(v) for k, v in (("Author", cp.author), ("Subject", cp.subject),
                                   ("Created", cp.created), ("Modified", cp.modified)) if v}
    return Extracted("docx", [(1, _clean("\n".join(lines)))], (cp.title or "").strip() or None, meta)


def _pptx(path: Path):
    from pptx import Presentation
    prs = Presentation(str(path))
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        parts = []
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        if title:
            parts.append(f"# {title}")
        for sh in slide.shapes:
            if sh.has_text_frame and sh != slide.shapes.title:
                t = sh.text_frame.text.strip()
                if t:
                    parts.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    parts.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        if parts:
            pages.append((i, _clean("\n\n".join(parts))))
    return Extracted("pptx", pages, (prs.core_properties.title or "").strip() or None,
                     {"slide_count": len(prs.slides)})


def _xlsx(path: Path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    pages = []
    for i, ws in enumerate(wb.worksheets, 1):
        rows = []
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v).strip() for v in row]
            while vals and not vals[-1]:
                vals.pop()
            if any(vals):
                rows.append("| " + " | ".join(vals) + " |")
        if rows:
            pages.append((i, f"# Sheet: {ws.title}\n\n" + "\n".join(rows)))
    wb.close()
    return Extracted("xlsx", pages, None, {"sheets": [ws.title for ws in wb.worksheets]})


def _csv(path: Path):
    text = _decode(path.read_bytes())
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = ["| " + " | ".join(c.strip() for c in r) + " |" for r in csv.reader(io.StringIO(text), dialect) if any(r)]
    return Extracted("csv", [(1, "\n".join(rows))], None, {"rows": len(rows)})


def _json(path: Path):
    raw = _decode(path.read_bytes())
    try:
        raw = json.dumps(json.loads(raw), indent=1, ensure_ascii=False)
    except ValueError:
        pass
    return Extracted("json", [(1, _clean(raw))])


class _Html(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "section", "article", "table", "ul", "ol", "pre", "blockquote"}
    SKIP = {"script", "style", "noscript", "svg", "head"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out, self.skip, self.title, self._in_title, self._heading = [], 0, None, False, None

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        if tag == "title":
            self._in_title = True
        if re.fullmatch(r"h[1-6]", tag):
            self._heading = int(tag[1])
            self.out.append("\n\n" + "#" * self._heading + " ")
        elif tag in self.BLOCK:
            self.out.append("\n")
        if tag == "li":
            self.out.append("- ")
        if tag in ("td", "th"):
            self.out.append(" | ")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        if tag == "title":
            self._in_title = False
        if re.fullmatch(r"h[1-6]", tag):
            self._heading = None
            self.out.append("\n\n")
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title = (self.title or "") + data.strip()
        elif not self.skip:
            self.out.append(re.sub(r"\s+", " ", data) if self._heading else data)


def _html(path: Path):
    p = _Html()
    p.feed(_decode(path.read_bytes()))
    text = re.sub(r"[ \t]+", " ", "".join(p.out))
    return Extracted("html", [(1, _clean(text))], p.title or None)


def _eml(path: Path):
    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    head = "\n".join(f"{h}: {msg[h]}" for h in ("From", "To", "Cc", "Date", "Subject") if msg[h])
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    if body is not None and body.get_content_type() == "text/html":
        p = _Html(); p.feed(text); text = "".join(p.out)
    attachments = [a.get_filename() for a in msg.iter_attachments() if a.get_filename()]
    if attachments:
        head += "\nAttachments: " + ", ".join(attachments)
    return Extracted("email", [(1, _clean(head + "\n\n" + text))], msg["Subject"],
                     {k: str(msg[k]) for k in ("From", "Date") if msg[k]})


EXTRACTORS = {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx, ".xlsm": _xlsx,
              ".csv": _csv, ".tsv": _csv, ".json": _json, ".html": _html, ".htm": _html, ".eml": _eml}
SUPPORTED = sorted(set(EXTRACTORS) | TEXT_EXT)


def supported(path: Path) -> bool:
    return path.suffix.lower() in EXTRACTORS or path.suffix.lower() in TEXT_EXT


def extract(path: Path) -> Extracted:
    ext = path.suffix.lower()
    if ext in EXTRACTORS:
        return EXTRACTORS[ext](path)
    if ext in TEXT_EXT:
        return _text(path)
    raise ValueError(f"unsupported file type {ext!r} (supported: {', '.join(SUPPORTED)})")


# ---- chunking ----------------------------------------------------------------------------
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")


def _hard_split(text, size):
    """Cut at line breaks, else spaces, else anywhere."""
    pieces = []
    while len(text) > size:
        cut = text.rfind("\n", 0, size)
        if cut < size // 2:
            cut = text.rfind(" ", 0, size)
        if cut < size // 2:
            cut = size
        pieces.append(text[:cut].strip())
        text = text[cut:].strip()
    return pieces + [text]


def chunk(pages, size, overlap):
    """Split pages into ~size-char chunks on paragraph boundaries, with `overlap` chars carried over.

    Returns [(page, heading, text)]. A chunk never spans pages, so citations point at one page.
    """
    out = []
    heading = None
    for page, text in pages:
        buf, buf_len, buf_heading = [], 0, heading

        def flush(carry=True):
            nonlocal buf, buf_len, buf_heading
            body = "\n\n".join(buf).strip()
            if body:
                out.append((page, buf_heading, body))
            tail = body[-overlap:] if carry and overlap and body else ""
            if tail and " " in tail:
                tail = tail[tail.index(" ") + 1:]   # start the overlap on a word boundary
            buf, buf_len, buf_heading = ([tail], len(tail), heading) if tail else ([], 0, heading)

        for para in re.split(r"\n\s*\n", text):
            para = para.strip()
            if not para:
                continue
            m = HEADING_RE.match(para.split("\n", 1)[0])
            if m:
                if buf_len > overlap:   # a new section starts a new chunk
                    flush(carry=False)
                heading = m.group(2).strip()[:200]
                if not buf or buf_heading is None:   # a chunk is labelled with its first section
                    buf_heading = heading
            if len(para) > size:   # one huge paragraph (tables, minified text): hard split
                pieces = _hard_split(para, size)
                para = pieces.pop()
                for piece in pieces:
                    if buf_len + len(piece) > size and buf:
                        flush()
                    buf.append(piece)
                    buf_len += len(piece)
                    flush()
            if buf_len + len(para) > size and buf:
                flush()
            buf.append(para)
            buf_len += len(para) + 2
        flush(carry=False)
    return out
